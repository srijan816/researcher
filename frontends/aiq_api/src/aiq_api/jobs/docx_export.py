# SPDX-FileCopyrightText: Copyright (c) 2025-2026, NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
# http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Markdown-to-DOCX conversion for the async report endpoint.

The markdown parsing helpers are dependency-free; only
:func:`markdown_to_docx_bytes` requires ``python-docx``, which is imported
lazily so deployments without the package degrade gracefully (the report
route returns HTTP 501).

Supported markdown:
- ``#``/``##``/``###`` (and deeper) headings
- paragraphs with ``**bold**``, ``*italic*``, and ```` `inline code` ````
- bullet (``-``/``*``/``+``) and numbered (``1.``/``1)``) lists
- links rendered as ``text (url)``
- horizontal rules rendered as spacing breaks
- fenced code blocks rendered as monospace paragraphs
"""

from __future__ import annotations

import re

DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_HR_RE = re.compile(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$")
_BULLET_RE = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUMBER_RE = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_FENCE_RE = re.compile(r"^\s*```")
_LINK_RE = re.compile(r"\[([^\]]+)\]\((\S+?)\)")
_INLINE_RE = re.compile(r"(\*\*[^*\n]+\*\*|\*[^*\n]+\*|`[^`\n]+`)")


def render_links_as_text(text: str) -> str:
    """Rewrite markdown links ``[text](url)`` as ``text (url)``."""
    return _LINK_RE.sub(lambda m: f"{m.group(1)} ({m.group(2)})", text)


def parse_inline(text: str) -> list[tuple[str, str]]:
    """Tokenize inline markdown into (style, text) runs.

    Styles: ``text``, ``bold``, ``italic``, ``code``. Links are rendered as
    plain ``text (url)`` before tokenization.
    """
    runs: list[tuple[str, str]] = []
    for part in _INLINE_RE.split(render_links_as_text(text)):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            runs.append(("bold", part[2:-2]))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            runs.append(("code", part[1:-1]))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            runs.append(("italic", part[1:-1]))
        else:
            runs.append(("text", part))
    return runs


def plain_text(text: str) -> str:
    """Strip inline markdown markers, keeping link targets as ``text (url)``."""
    return "".join(run_text for _style, run_text in parse_inline(text))


def parse_markdown_blocks(markdown: str) -> list[tuple]:
    """Parse markdown into a flat list of block tuples.

    Block shapes:
    - ``("heading", level, text)``
    - ``("paragraph", text)``
    - ``("bullet", text)`` / ``("number", text)``
    - ``("code", text)``
    - ``("hr",)``
    """
    blocks: list[tuple] = []
    paragraph_lines: list[str] = []
    code_lines: list[str] | None = None

    def flush_paragraph() -> None:
        nonlocal paragraph_lines
        if paragraph_lines:
            blocks.append(("paragraph", " ".join(line.strip() for line in paragraph_lines)))
            paragraph_lines = []

    for line in (markdown or "").splitlines():
        if code_lines is not None:
            if _FENCE_RE.match(line):
                blocks.append(("code", "\n".join(code_lines)))
                code_lines = None
            else:
                code_lines.append(line)
            continue

        if _FENCE_RE.match(line):
            flush_paragraph()
            code_lines = []
            continue

        if not line.strip():
            flush_paragraph()
            continue

        heading = _HEADING_RE.match(line)
        if heading:
            flush_paragraph()
            blocks.append(("heading", len(heading.group(1)), heading.group(2).strip()))
            continue

        if _HR_RE.match(line):
            flush_paragraph()
            blocks.append(("hr",))
            continue

        bullet = _BULLET_RE.match(line)
        if bullet:
            flush_paragraph()
            blocks.append(("bullet", bullet.group(1).strip()))
            continue

        number = _NUMBER_RE.match(line)
        if number:
            flush_paragraph()
            blocks.append(("number", number.group(1).strip()))
            continue

        paragraph_lines.append(line)

    if code_lines is not None:
        blocks.append(("code", "\n".join(code_lines)))
    flush_paragraph()
    return blocks


def _add_runs(paragraph, text: str) -> None:
    """Add styled runs for inline markdown to a python-docx paragraph."""
    for style, run_text in parse_inline(text):
        run = paragraph.add_run(run_text)
        if style == "bold":
            run.bold = True
        elif style == "italic":
            run.italic = True
        elif style == "code":
            run.font.name = "Courier New"


def markdown_to_docx_bytes(markdown: str, *, title: str | None = None) -> bytes:
    """Convert report markdown into a .docx file.

    Raises:
        ImportError: when ``python-docx`` is not installed. Callers should map
            this to HTTP 501 with an install hint.
    """
    import io

    import docx  # Lazy: optional dependency (python-docx)
    from docx.shared import Pt

    document = docx.Document()
    if title:
        document.core_properties.title = title

    for block in parse_markdown_blocks(markdown):
        kind = block[0]
        if kind == "heading":
            level = min(int(block[1]), 4)
            document.add_heading(plain_text(block[2]), level=level)
        elif kind == "paragraph":
            paragraph = document.add_paragraph()
            _add_runs(paragraph, block[1])
        elif kind == "bullet":
            paragraph = document.add_paragraph(style="List Bullet")
            _add_runs(paragraph, block[1])
        elif kind == "number":
            paragraph = document.add_paragraph(style="List Number")
            _add_runs(paragraph, block[1])
        elif kind == "code":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block[1])
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        elif kind == "hr":
            # Page-ish break: extra empty paragraph keeps sections visually separated
            # without exploding page count for reports that use --- separators.
            document.add_paragraph()

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
